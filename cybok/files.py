"""Secure file storage and text extraction for uploaded documents."""
import csv
import io
import json
import re
import statistics
import zipfile
from pathlib import Path

from flask import current_app
from PIL import Image, UnidentifiedImageError

from .extensions import db
from .models import File, new_id

MAX_TEXT = 300_000

DOC_EXT = {"pdf", "docx", "pptx", "odt", "rtf", "txt", "md", "markdown", "html", "htm", "xml", "json", "log"}
DATA_EXT = {"csv", "tsv", "xlsx", "xlsm"}
CODE_EXT = {"py", "js", "ts", "jsx", "tsx", "java", "c", "h", "cpp", "hpp", "cs", "go", "rs", "rb", "php", "sql",
            "yaml", "yml", "ini", "toml", "css", "scss", "sh", "bat", "ps1", "kt", "swift", "dart", "r", "m", "lua"}
IMAGE_EXT = {"png", "jpg", "jpeg", "gif", "webp", "bmp"}
VIDEO_EXT = {"mp4", "webm"}  # generated only, not accepted as uploads
ALLOWED_EXT = DOC_EXT | DATA_EXT | CODE_EXT | IMAGE_EXT

MIME = {
    "pdf": "application/pdf", "txt": "text/plain", "md": "text/markdown", "csv": "text/csv",
    "json": "application/json", "png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
    "gif": "image/gif", "webp": "image/webp", "bmp": "image/bmp",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


class UploadError(Exception):
    pass


def storage_root():
    root = Path(current_app.config["UPLOAD_DIR"])
    root.mkdir(parents=True, exist_ok=True)
    return root


def abs_path(rel):
    root = storage_root().resolve()
    p = (root / rel).resolve()
    if root not in p.parents and p != root:
        raise UploadError("Invalid storage path.")
    return p


def file_ext(filename):
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def safe_display_name(filename):
    name = re.sub(r"[\x00-\x1f<>:\"/\\|?*]", "_", (filename or "file").strip())[:200]
    return name or "file"


def kind_for(ext):
    if ext in IMAGE_EXT:
        return "image"
    if ext in VIDEO_EXT:
        return "video"
    if ext in DATA_EXT:
        return "data"
    if ext in CODE_EXT:
        return "code"
    return "document"


def _check_magic(ext, data):
    if ext == "pdf" and not data.startswith(b"%PDF"):
        raise UploadError("This file doesn't look like a valid PDF.")
    if ext in {"docx", "xlsx", "xlsm", "pptx", "odt"} and not data.startswith(b"PK"):
        raise UploadError(f"This file doesn't look like a valid .{ext} document.")


def save_bytes(user_id, data, filename, mime=None, source="upload", meta=None, in_library=True, folder=""):
    """Stores bytes under a random name in the user's private folder and creates the File row."""
    ext = file_ext(filename)
    fid = new_id()
    rel = f"{user_id}/{fid}.{ext or 'bin'}"
    path = abs_path(rel)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    meta = dict(meta or {})
    kind = kind_for(ext)
    if kind == "image":
        try:
            with Image.open(io.BytesIO(data)) as im:
                meta.update(width=im.width, height=im.height)
                thumb = im.convert("RGB") if im.mode not in ("RGB", "RGBA") else im.copy()
                thumb.thumbnail((480, 480))
                thumb.save(str(path) + ".thumb.webp", "WEBP", quality=80)
        except (UnidentifiedImageError, OSError):
            path.unlink(missing_ok=True)
            raise UploadError("This image file is damaged or not supported.")
    f = File(id=fid, user_id=user_id, name=safe_display_name(filename), mime=mime or MIME.get(ext, "application/octet-stream"),
             ext=ext, size=len(data), storage_path=rel, source=source, kind=kind, meta=json.dumps(meta),
             in_library=in_library, folder=folder)
    db.session.add(f)
    return f


def handle_upload(user_id, filename, data, max_mb):
    filename = safe_display_name(filename or "file")
    ext = file_ext(filename)
    if ext not in ALLOWED_EXT:
        raise UploadError(f"'.{ext or '?'}' files aren't supported. Try PDF, DOCX, TXT, CSV, XLSX, PPTX, code or images.")
    if len(data) > max_mb * 1024 * 1024:
        raise UploadError(f"File is larger than {max_mb} MB.")
    if not data:
        raise UploadError("The file is empty.")
    _check_magic(ext, data)
    f = save_bytes(user_id, data, filename)
    try:
        text, extra = extract_text(ext, data)
    except Exception as exc:  # extraction is best-effort
        text, extra = None, {"extract_error": str(exc)[:200]}
    if text:
        f.extracted_text = text[:MAX_TEXT]
    if extra:
        meta = f.meta_dict
        meta.update(extra)
        f.meta = json.dumps(meta)
    return f


def delete_file_blob(f):
    try:
        p = abs_path(f.storage_path)
        p.unlink(missing_ok=True)
        Path(str(p) + ".thumb.webp").unlink(missing_ok=True)
    except (UploadError, OSError):
        pass


# ------------------------------------------------------------------ extraction
def _decode(data):
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            if enc == "utf-16" and "\x00" in text:
                continue
            return text
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf(data):
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        if t.strip():
            pages.append(f"--- Page {i} ---\n{t.strip()}")
        if sum(len(p) for p in pages) > MAX_TEXT:
            break
    text = "\n\n".join(pages)
    extra = {"pages": len(reader.pages)}
    if not text.strip():
        extra["note"] = "No selectable text found (the PDF may be scanned images)."
    return text, extra


def _docx(data):
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = []
    for p in d.paragraphs:
        if p.text.strip():
            style = (p.style.name or "").lower() if p.style is not None else ""
            prefix = "#" * int(style[-1]) + " " if style.startswith("heading") and style[-1:].isdigit() else ""
            parts.append(prefix + p.text.strip())
    for ti, table in enumerate(d.tables, 1):
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        if rows:
            parts.append(f"\n[Table {ti}]\n" + "\n".join(rows))
    return "\n".join(parts), {"paragraphs": len(d.paragraphs), "tables": len(d.tables)}


def _zip_xml_text(data, prefix, tag_regex):
    out = []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = sorted((n for n in z.namelist() if n.startswith(prefix) and n.endswith(".xml")),
                       key=lambda n: int(re.findall(r"\d+", n)[-1]) if re.findall(r"\d+", n) else 0)
        for i, n in enumerate(names, 1):
            xml = z.read(n).decode("utf-8", errors="ignore")
            texts = [re.sub(r"<[^>]+>", "", t) for t in re.findall(tag_regex, xml, flags=re.S)]
            joined = " ".join(t for t in texts if t.strip())
            if joined.strip():
                out.append((i, joined))
    return out


def _pptx(data):
    slides = _zip_xml_text(data, "ppt/slides/slide", r"<a:t>(.*?)</a:t>")
    return "\n\n".join(f"--- Slide {i} ---\n{t}" for i, t in slides), {"slides": len(slides)}


def _odt(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read("content.xml").decode("utf-8", errors="ignore")
    xml = re.sub(r"</text:p>|</text:h>", "\n", xml)
    return re.sub(r"<[^>]+>", "", xml), {}


def _html(text):
    text = re.sub(r"(?is)<(script|style).*?</\1>", "", text)
    text = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</h\d>|</li>", "\n", text)
    text = re.sub(r"<[^>]+>", "", text)
    return re.sub(r"\n{3,}", "\n\n", text)


def _table_summary(sheets):
    """sheets: list of (name, rows). Produces a compact, model-friendly description with statistics."""
    out, meta = [], {"sheets": []}
    for name, rows in sheets:
        rows = [r for r in rows if any(str(c).strip() for c in r if c is not None)]
        if not rows:
            continue
        header = [str(c).strip() if c is not None else f"col{i + 1}" for i, c in enumerate(rows[0])]
        body = rows[1:]
        meta["sheets"].append({"name": name, "rows": len(body), "columns": len(header)})
        out.append(f"## Sheet: {name}\nRows: {len(body)} | Columns: {len(header)}\nColumns: {', '.join(header)}")
        stats = []
        for ci, col in enumerate(header):
            nums = []
            for r in body:
                if ci < len(r) and r[ci] is not None and str(r[ci]).strip() != "":
                    try:
                        nums.append(float(str(r[ci]).replace(",", "")))
                    except ValueError:
                        pass
            if len(nums) >= max(2, int(len(body) * 0.6)):
                stats.append(f"- {col}: count={len(nums)}, sum={sum(nums):,.4g}, mean={statistics.fmean(nums):,.4g}, "
                             f"min={min(nums):,.4g}, max={max(nums):,.4g}")
        if stats:
            out.append("Numeric column statistics:\n" + "\n".join(stats))
        limit = 400
        lines = [" | ".join(header)]
        for r in body[:limit]:
            lines.append(" | ".join("" if c is None else str(c) for c in r))
        more = f"\n... ({len(body) - limit} more rows not shown)" if len(body) > limit else ""
        out.append("Data:\n" + "\n".join(lines) + more)
    return "\n\n".join(out), meta


def _csv(data, delimiter=None):
    text = _decode(data)
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    return _table_summary([("data", rows)])


def _xlsx(data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets[:10]:
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > 5000:
                break
            rows.append(list(row))
        sheets.append((ws.title, rows))
    wb.close()
    return _table_summary(sheets)


def extract_text(ext, data):
    if ext == "pdf":
        return _pdf(data)
    if ext == "docx":
        return _docx(data)
    if ext == "pptx":
        return _pptx(data)
    if ext == "odt":
        return _odt(data)
    if ext in {"xlsx", "xlsm"}:
        return _xlsx(data)
    if ext == "csv":
        return _csv(data)
    if ext == "tsv":
        return _csv(data, "\t")
    if ext in {"html", "htm"}:
        return _html(_decode(data)), {}
    if ext == "rtf":
        t = _decode(data)
        t = re.sub(r"\\[a-z]+-?\d* ?|[{}]", "", t)
        return t, {}
    if ext in IMAGE_EXT:
        return None, {}
    text = _decode(data)
    return text, {"lines": text.count("\n") + 1}
